import shutil

import fitz
import openpyxl
import pytest

from verideck.pipeline import UPLOADS, load_results, run_job

needs_libreoffice = pytest.mark.skipif(
    shutil.which("soffice") is None and shutil.which("libreoffice") is None,
    reason="LibreOffice not installed",
)


def make_job(tmp_path):
    """A deck (PDF) and a model (XLSX) sharing one value and disagreeing on another."""
    uploads = tmp_path / UPLOADS
    uploads.mkdir(parents=True)

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Total sales 1,234.56")
    page.insert_text((72, 130), "Headcount cost 500.25")
    doc.save(uploads / "deck.pdf")
    doc.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Total sales", 1234.56
    ws["A2"], ws["B2"] = "Headcount cost", 500.75  # does NOT tie with the deck
    # Wide enough that the label is not clipped in the PDF rendering; a
    # clipped label cannot be matched (known limitation of rendering-based
    # extraction).
    ws.column_dimensions["A"].width = 24
    wb.save(uploads / "model.xlsx")
    return tmp_path


@needs_libreoffice
def test_run_job_end_to_end(tmp_path):
    job_dir = make_job(tmp_path)

    results = run_job(job_dir)

    assert [f["name"] for f in results["files"]] == ["deck.pdf", "model.xlsx"]
    ties = [g for g in results["groups"] if g["kind"] == "tie"]
    mismatches = [g for g in results["groups"] if g["kind"] == "mismatch"]

    tie_keys = {g["key"] for g in ties}
    assert "1234.56" in tie_keys, f"expected cross-file tie, got {tie_keys}"
    sales = next(g for g in ties if g["key"] == "1234.56")
    assert {o["file"] for o in sales["occurrences"]} == {"deck.pdf", "model.xlsx"}

    assert any(g["key"] == "headcount cost" for g in mismatches)

    for group in results["groups"]:
        for occurrence in group["occurrences"]:
            crop = job_dir / "crops" / occurrence["crop"]
            assert crop.exists() and crop.stat().st_size > 0

    assert load_results(job_dir) == results
