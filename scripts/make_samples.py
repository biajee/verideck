"""Generate a sample reporting package in samples/ for demoing Verideck.

The package describes one fictional quarter (Q4 FY2025) across four files:

  board_deck.pptx   4 slides
  annual_report.pdf 3 pages
  model.xlsx        the "source of truth" spreadsheet
  scanned_memo.pdf  image-only page, exercises OCR

Built-in storylines to look for on the results page:
  * Revenue 12,456,789 ties across every document, including the scan.
  * COGS appears as (7,234,512) in the PDF and -7,234,512 in Excel — same value.
  * EPS is shown as 1.17 in the deck/memo but 1.166 in the PDF/model:
    a tie within rounding (amber badge).
  * EBITDA is 2,076,654 in the deck and PDF but 2,067,654 in Excel —
    a deliberate transposition error Verideck must flag for review.

Run: .venv/bin/python scripts/make_samples.py
"""

from pathlib import Path

import fitz
import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def add_slide(prs, title, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.5), Inches(5.5))
    frame = box.text_frame
    frame.text = title
    frame.paragraphs[0].font.size = Pt(32)
    for line in lines:
        para = frame.add_paragraph()
        para.text = line
        para.font.size = Pt(22)


def make_pptx(path: Path) -> None:
    prs = Presentation()
    add_slide(prs, "Q4 FY2025 Board Deck", [
        "Acme Holdings, Inc.",
        "Revenue $12,456,789",
    ])
    add_slide(prs, "P&L summary", [
        "Revenue 12,456,789",
        "Gross profit 5,222,277",
        "EBITDA 2,076,654",
        "Net income 1,456,098",
    ])
    add_slide(prs, "Cash & people", [
        "Cash balance 8,901,234",
        "Headcount 1,250",
    ])
    add_slide(prs, "Per-share highlights", [
        "EPS of 1.17 this quarter",  # rounded display of 1.166
    ])
    prs.save(path)


def add_pdf_page(doc, title, lines):
    page = doc.new_page()
    page.insert_text((72, 90), title, fontsize=16)
    for i, line in enumerate(lines):
        page.insert_text((72, 140 + 26 * i), line, fontsize=12)


def make_pdf(path: Path) -> None:
    doc = fitz.open()
    add_pdf_page(doc, "Annual Report — Income statement", [
        "Revenue 12,456,789",
        "Cost of goods sold (7,234,512)",
        "Gross profit 5,222,277",
        "Operating expenses (3,145,623)",
        "EBITDA 2,076,654",
        "Net income 1,456,098",
    ])
    add_pdf_page(doc, "Balance sheet extract", [
        "Cash balance 8,901,234",
        "EPS 1.166",
    ])
    add_pdf_page(doc, "Management discussion", [
        "Revenue of $12,456,789 was driven by strong renewals,",
        "with a headcount of 1,250 at quarter end.",
    ])
    doc.save(path)
    doc.close()


def make_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    rows = [
        ("Revenue", 12456789, "#,##0"),
        ("Cost of goods sold", -7234512, "#,##0"),
        ("Gross profit", 5222277, "#,##0"),
        ("Operating expenses", -3145623, "#,##0"),
        ("EBITDA", 2067654, "#,##0"),  # deliberate error: deck/PDF say 2,076,654
        ("Net income", 1456098, "#,##0"),
        ("Cash balance", 8901234, "#,##0"),
        ("EPS", 1.166, "0.000"),
        ("Headcount", 1250, "#,##0"),
    ]
    for row_no, (label, value, fmt) in enumerate(rows, start=1):
        ws.cell(row=row_no, column=1, value=label)
        cell = ws.cell(row=row_no, column=2, value=value)
        cell.number_format = fmt
    # Wide columns so the rendering shows full values: a narrow column makes
    # the PDF rendition display rounded numbers, which Verideck would then —
    # correctly — report as not tying with the other documents.
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 18
    wb.save(path)


def make_scanned_pdf(path: Path) -> None:
    """An image-only 'scan' (rasterized text), so reading it requires OCR."""
    src = fitz.open()
    page = src.new_page()
    page.insert_text((72, 100), "Scanned board memo", fontsize=16)
    page.insert_text((72, 150), "Revenue 12,456,789 confirmed by the committee.", fontsize=12)
    page.insert_text((72, 176), "EPS 1.17", fontsize=12)
    pix = page.get_pixmap(matrix=fitz.Matrix(3, 3))
    scan = fitz.open()
    scan_page = scan.new_page(width=page.rect.width, height=page.rect.height)
    # JPEG like a real scanner, else the embedded raster is ~13 MB raw
    scan_page.insert_image(scan_page.rect, stream=pix.tobytes("jpeg"))
    scan.save(path)
    scan.close()
    src.close()


def main() -> None:
    SAMPLES.mkdir(exist_ok=True)
    make_pptx(SAMPLES / "board_deck.pptx")
    make_pdf(SAMPLES / "annual_report.pdf")
    make_xlsx(SAMPLES / "model.xlsx")
    make_scanned_pdf(SAMPLES / "scanned_memo.pdf")
    print(f"Samples written to {SAMPLES}")


if __name__ == "__main__":
    main()
